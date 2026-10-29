#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
核心资源标记表加载与匹配工具。
读取"关键资源标记表.xlsx"，提取核心业务资源的 resource_id 集合，
用于在分析时判断告警是否属于核心资源。
按 Region 列过滤，支持多区域。
"""

import os


def load_core_regions(xlsx_path: str) -> list:
    """读取关键资源标记表.xlsx，返回核心业务资源涉及的 Region 列表。

    筛选"是否核心业务"列值为"是"的行，取"Region"列去重。
    返回 list[str]，保持出现顺序。
    """
    if not xlsx_path or not os.path.exists(xlsx_path):
        print(f"[WARN] 核心资源标记表不存在: {xlsx_path}")
        return []

    try:
        import openpyxl
    except ImportError:
        print("[ERROR] 读取 xlsx 需要 openpyxl，请运行: pip install openpyxl")
        return []

    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    ws = wb[wb.sheetnames[0]]

    regions = []
    seen = set()
    header = {}
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            for j, cell in enumerate(row):
                if cell is not None:
                    header[str(cell).strip()] = j
            continue

        core_col = header.get('是否核心业务')
        region_col = header.get('Region')
        if core_col is None or region_col is None:
            continue

        is_core = row[core_col] if core_col < len(row) else None
        region_val = row[region_col] if region_col < len(row) else None

        if not (is_core and str(is_core).strip() == '是' and region_val):
            continue

        region_str = str(region_val).strip()
        if region_str not in seen:
            seen.add(region_str)
            regions.append(region_str)

    wb.close()
    print(f"Loaded {len(regions)} regions from {os.path.basename(xlsx_path)}: {regions}")
    return regions


def load_core_resources(xlsx_path: str, region: str = None):
    """读取关键资源标记表.xlsx，返回核心资源 resource_id 集合。

    筛选"是否核心业务"列值为"是"的行，取"resource_id"列值。
    若提供 region，则额外按 Excel 中"Region"列过滤，返回 set[str]。
    若 region 为 None，则返回 {region: set(resource_ids)} 字典供按区域筛选。
    """
    if not xlsx_path or not os.path.exists(xlsx_path):
        print(f"[WARN] 核心资源标记表不存在: {xlsx_path}")
        return set() if region else {}

    try:
        import openpyxl
    except ImportError:
        print("[ERROR] 读取 xlsx 需要 openpyxl，请运行: pip install openpyxl")
        return set() if region else {}

    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    ws = wb[wb.sheetnames[0]]

    # 按区域返回字典
    if region is None:
        region_map = {}
        header = {}
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                for j, cell in enumerate(row):
                    if cell is not None:
                        header[str(cell).strip()] = j
                continue

            core_col = header.get('是否核心业务')
            rid_col = header.get('resource_id')
            region_col = header.get('Region')
            if core_col is None or rid_col is None or region_col is None:
                continue

            is_core = row[core_col] if core_col < len(row) else None
            resource_id = row[rid_col] if rid_col < len(row) else None
            region_val = row[region_col] if region_col < len(row) else None

            if not (is_core and str(is_core).strip() == '是' and resource_id and region_val):
                continue

            region_str = str(region_val).strip()
            if region_str not in region_map:
                region_map[region_str] = set()
            region_map[region_str].add(str(resource_id).strip())

        wb.close()
        total = sum(len(v) for v in region_map.values())
        print(f"Loaded {total} core resources across {len(region_map)} regions from {os.path.basename(xlsx_path)}")
        return region_map

    # 按指定区域返回 set
    core_ids = set()
    header = {}
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            for j, cell in enumerate(row):
                if cell is not None:
                    header[str(cell).strip()] = j
            continue

        core_col = header.get('是否核心业务')
        rid_col = header.get('resource_id')
        region_col = header.get('Region')
        if core_col is None or rid_col is None or region_col is None:
            continue

        is_core = row[core_col] if core_col < len(row) else None
        resource_id = row[rid_col] if rid_col < len(row) else None
        region_val = row[region_col] if region_col < len(row) else None

        if not (is_core and str(is_core).strip() == '是' and resource_id and region_val):
            continue

        if str(region_val).strip() != region:
            continue

        core_ids.add(str(resource_id).strip())

    wb.close()
    print(f"Loaded {len(core_ids)} core resources for region {region} from {os.path.basename(xlsx_path)}")
    return core_ids


def is_core_resource(alarm: dict, core_ids: set) -> bool:
    """判断告警的 resourceId 是否在核心资源集合中。

    仅用 resource_id 精确匹配（UUID 对 UUID）。
    core_ids 为空时返回 False（未加载标记表 = 不标记）。
    """
    if not core_ids:
        return False
    res = alarm.get('alarm_resource') or {}
    resource_id = res.get('resourceId', '')
    return bool(resource_id and resource_id in core_ids)
