# AI-generated
#!/usr/bin/env python3
"""Generate eval fixtures (repeatable, deterministic).

Fixtures are intentionally small: fast to run in regression and easy to see which boundary
a failure hit.
- no_other_broker/    one broker (many disks with space): 0 distinct-broker placements, and
                      same-broker different-mount cannot be expressed in Kafka reassignment
                      JSON -> NO_OTHER_BROKER block is certain
- l3_block/           one broker, one disk already at 99% -> L3 block is certain
- pre_existing_over_threshold/  an input disk is already over threshold; validation should
                      WARN instead of failing the stage

Usage: python evals/build_fixtures.py
"""

import os
import shutil
import sys

sys.dont_write_bytecode = True

import openpyxl

EVAL_DIR = os.path.dirname(os.path.abspath(__file__))
FIXTURES = os.path.join(EVAL_DIR, 'fixtures')
ASSETS = os.path.join(os.path.dirname(EVAL_DIR), 'assets')

TOPIC_HEADERS = ['Topic', 'PartitionCount', 'Partition', 'Broker ID', 'Path',
                 'Size(bytes)', 'Size(GB)']
BROKER_HEADERS = ['Broker ID', 'Host', 'Disk (Used | Total)', '分区', '挂载点',
                  '已使用（GB）', '磁盘大小（GB）', '磁盘使用率']

DISK_GB = 922.21   # same order of magnitude as the assets template, easier to compare


def _write_topic(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Topic'
    ws.append(TOPIC_HEADERS)
    for r in rows:
        ws.append(r)
    wb.save(path)
    wb.close()


def _write_broker(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Broker'
    ws.append(BROKER_HEADERS)
    for r in rows:
        ws.append(r)
    wb.save(path)
    wb.close()


def _broker_row(bid, host, mnt, used_gb):
    usage = round(used_gb / DISK_GB * 100, 2)
    return [bid, host, '%.1fTB | %.1fTB' % (used_gb / 1024.0, DISK_GB / 1024.0), '/dev/vdb1',
            mnt, used_gb, DISK_GB, usage]


def build_no_other_broker():
    d = os.path.join(FIXTURES, 'no_other_broker')
    os.makedirs(d, exist_ok=True)
    topic_rows = []
    for p in range(3):
        topic_rows.append(['ITDE_TEST_single_broker_lkp', 3, p, 1, '/srv/BigData/data1',
                           107374182 * (p + 1), round(0.1 * (p + 1), 1)])
    _write_topic(os.path.join(d, 'topic.xlsx'), topic_rows)
    _write_broker(os.path.join(d, 'broker.xlsx'), [
        _broker_row(1, '10.51.128.165', '/srv/BigData/data%d' % i, 92.0 + i)
        for i in range(1, 6)
    ])


def build_l3():
    d = os.path.join(FIXTURES, 'l3_block')
    os.makedirs(d, exist_ok=True)
    topic_rows = []
    for p in range(2):
        topic_rows.append(['ITDE_TEST_no_slot_lkp', 2, p, 1, '/srv/BigData/data1',
                           107374182 * (p + 1), round(0.1 * (p + 1), 1)])
    _write_topic(os.path.join(d, 'topic.xlsx'), topic_rows)
    # The only disk is already at 99%; any new replica pushes the projection above the threshold -> must block
    _write_broker(os.path.join(d, 'broker.xlsx'),
                  [_broker_row(1, '10.51.128.165', '/srv/BigData/data1', DISK_GB * 0.99)])


def build_pre_existing():
    d = os.path.join(FIXTURES, 'pre_existing_over_threshold')
    os.makedirs(d, exist_ok=True)
    src = os.path.join(ASSETS, 'broker_template.xlsx')
    wb = openpyxl.load_workbook(src)
    ws = wb.active
    for row in ws.iter_rows(min_row=2):
        if str(row[0].value) == '1' and str(row[4].value) == '/srv/BigData/data1':
            row[5].value = 830
            row[7].value = 90.0
    wb.save(os.path.join(d, 'broker.xlsx'))
    wb.close()
    shutil.copy(os.path.join(ASSETS, 'topic_template.xlsx'), os.path.join(d, 'topic.xlsx'))


def main():
    os.makedirs(FIXTURES, exist_ok=True)
    build_no_other_broker()
    build_l3()
    build_pre_existing()
    for root, _dirs, files in os.walk(FIXTURES):
        for f in sorted(files):
            p = os.path.join(root, f)
            print('%-70s %d bytes' % (os.path.relpath(p, EVAL_DIR), os.path.getsize(p)))


if __name__ == '__main__':
    main()
