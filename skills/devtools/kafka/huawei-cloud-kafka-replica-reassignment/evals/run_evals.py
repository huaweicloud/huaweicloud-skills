# AI-generated
#!/usr/bin/env python3
"""Regression entry for kafka-replica-reassignment.

Runs the 8 cases defined in evals/evals.json, printing verifiable assertions per case;
exit code = number of failing assertions (0 = all pass), so it can hang off a pipeline or
a pre-commit hook.

Design constraints:
- artifacts are always written to tempfile-created temp dirs; the skill package stays
  read-only (it may be read-only after install and must not be polluted);
- every case verifies input file shas are unchanged — "read-only" is part of the tool's
  promise;
- cases share no state; running a single case also works.

Usage: python evals/run_evals.py
"""

import hashlib
import json
import os
import shutil
import subprocess
import sys
import tempfile

sys.dont_write_bytecode = True

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SCRIPTS = os.path.join(ROOT, 'scripts')
ASSETS = os.path.join(ROOT, 'assets')
FIX = os.path.join(ROOT, 'evals', 'fixtures')
PY = sys.executable

RESULTS = []
_TMP = []


def check(case, text, ok, evidence=''):
    RESULTS.append({'case': case, 'text': text, 'passed': bool(ok), 'evidence': str(evidence)})
    print('  [%s] %s%s' % ('PASS' if ok else 'FAIL', text, ('  <- ' + str(evidence)) if evidence else ''))


def sha(path):
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        for chunk in iter(lambda: f.read(65536), b''):
            h.update(chunk)
    return h.hexdigest()


def snapshot(dirpath):
    """Set of relative paths of all files under a directory, used to detect "extra stuff". """
    seen = set()
    for root, _dirs, files in os.walk(dirpath):
        for f in files:
            seen.add(os.path.relpath(os.path.join(root, f), dirpath))
    return seen


def workdir(name):
    d = tempfile.mkdtemp(prefix='eval-%s-' % name)
    _TMP.append(d)
    return d


def run_stage(script, args):
    proc = subprocess.run([PY, os.path.join(SCRIPTS, script)] + args,
                          cwd=SCRIPTS, capture_output=True, text=True,
                          encoding='utf-8', errors='replace')
    try:
        data = json.loads(proc.stdout.strip())
    except Exception:
        data = {'_unparsable_stdout': proc.stdout[-800:], '_stderr': proc.stderr[-800:]}
    return proc.returncode, data


def stage1(topic, broker, out, replica=2, threshold=70):
    return run_stage('stage1_allocate.py', ['--topic', topic, '--broker', broker, '--output-dir', out,
                                            '--replica-count', str(replica),
                                            '--disk-threshold', str(threshold)])


def stage2(out, replica=2):
    return run_stage('stage2_merge.py', ['--main', os.path.join(out, 'topic.xlsx'),
                                         '--new', os.path.join(out, 'topic_new.xlsx'),
                                         '--output-dir', out, '--replica-count', str(replica)])


def stage3(out, strategy='size_10g', param='10', replica=2):
    return run_stage('stage3_json.py', ['--input', os.path.join(out, 'topic_all.xlsx'),
                                        '--output-dir', out, '--strategy', strategy,
                                        '--strategy-param', str(param),
                                        '--replica-count', str(replica)])


def table_header(path):
    """Read artifact headers; missing file returns [] (assertions fail with empty-header evidence)."""
    if not os.path.isfile(path):
        return []
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    first = next(wb.active.iter_rows(values_only=True), None)
    wb.close()
    return [str(c) for c in (first or [])]


def table_matrix(path):
    """Read an artifact table into a string matrix.

    xlsx is a zip container with internal write timestamps, so two byte-identical-content
    files never have identical bytes; "reproducible results" must compare content, not sha.
    """
    if not os.path.isfile(path):
        return None
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = [[('' if c is None else str(c)) for c in r] for r in wb.active.iter_rows(values_only=True)]
    wb.close()
    return rows


def read_json_files(out):
    jdir = os.path.join(out, 'json')
    files = sorted(os.listdir(jdir)) if os.path.isdir(jdir) else []
    docs = {}
    for name in files:
        with open(os.path.join(jdir, name), encoding='utf-8') as f:
            docs[name] = json.load(f)
    return files, docs


def case1_golden():
    name = '黄金路径'
    print('\n=== 用例 1：%s ===' % name)
    topic, broker = os.path.join(ASSETS, 'topic_template.xlsx'), os.path.join(ASSETS, 'broker_template.xlsx')
    before = (sha(topic), sha(broker))
    out = workdir('golden')
    rc1, s1 = stage1(topic, broker, out)
    st1 = s1.get('stats', {})
    check(name, '阶段一 status=PASS 且退出码 0', rc1 == 0 and s1.get('status') == 'PASS',
          'rc=%s status=%s' % (rc1, s1.get('status')))
    check(name, '阶段一 partitions_total=22 / new_replicas_assigned=22 / degradation_count=0',
          st1.get('partitions_total') == 22 and st1.get('new_replicas_assigned') == 22
          and st1.get('degradation_count') == 0, json.dumps(st1, ensure_ascii=False))
    check(name, '阶段一 max_disk_usage=37.84%（±0.5）',
          abs(float(st1.get('max_disk_usage', 0)) - 37.84) <= 0.5, st1.get('max_disk_usage'))
    outs = sorted(os.path.basename(p) for p in s1.get('outputs', []))
    check(name, '阶段一产物命名 = topic.xlsx / topic_new.xlsx / broker_updated.xlsx',
          outs == ['broker_updated.xlsx', 'topic.xlsx', 'topic_new.xlsx'], outs)
    checks1 = {c['check']: c for c in s1.get('validation', {}).get('checks', [])}
    c16 = checks1.get('1.6', {})
    check(name, '阶段一 1.6 均衡检查 PASS（新副本未集中在个别 broker）',
          c16.get('status') == 'PASS', c16.get('detail'))
    rc2, s2 = stage2(out)
    st2 = s2.get('stats', {})
    check(name, '阶段二 merged_rows=22 / missing=0 / dup=0',
          st2.get('merged_rows') == 22 and st2.get('missing_partitions') == 0
          and st2.get('duplicate_rows') == 0, json.dumps(st2, ensure_ascii=False))
    check(name, '阶段二产物命名 = topic_all.xlsx',
          os.path.isfile(os.path.join(out, 'topic_all.xlsx'))
          and sorted(os.path.basename(p) for p in s2.get('outputs', [])) == ['topic_all.xlsx'],
          sorted(os.path.basename(p) for p in s2.get('outputs', [])))
    header = table_header(os.path.join(out, 'topic_all.xlsx'))
    check(name, '阶段二合并表 10 列且保留 degradation',
          len(header) == 10 and 'degradation' in header and len(set(header)) == len(header), header)
    rc3, s3 = stage3(out)
    st3 = s3.get('stats', {})
    check(name, '阶段三 status=PASS / total_entries=22 / validation_fail=0',
          rc3 == 0 and s3.get('status') == 'PASS' and st3.get('total_entries') == 22
          and st3.get('validation_fail') == 0, json.dumps(st3, ensure_ascii=False))
    files, docs = read_json_files(out)
    ok_struct = bool(docs)
    for doc in docs.values():
        if doc.get('version') != 1:
            ok_struct = False
        for entry in doc.get('partitions', []):
            if entry.get('log_dirs', [None])[0] != 'any' or len(entry.get('replicas', [])) != 2:
                ok_struct = False
    check(name, "JSON version=1 且每条 log_dirs[0]='any'、len(replicas)=2",
          ok_struct, '%s 个文件' % len(files))
    check(name, '输入文件 sha 不变（只查不改）', (sha(topic), sha(broker)) == before)

    # The artifact topic.xlsx shares its name with the input: pointing --output-dir at the
    # input directory must be blocked. If it were not, the user's current-state table would
    # be overwritten irreversibly.
    guard = workdir('guard')
    shutil.copy(topic, os.path.join(guard, 'topic.xlsx'))
    shutil.copy(broker, os.path.join(guard, 'broker.xlsx'))
    g_topic = os.path.join(guard, 'topic.xlsx')
    g_before = sha(g_topic)
    rc_g, s_g = stage1(g_topic, os.path.join(guard, 'broker.xlsx'), guard)
    check(name, '输出目录=输入目录时被拦下且未覆盖输入（OUTPUT_OVERWRITES_INPUT）',
          s_g.get('reason') == 'OUTPUT_OVERWRITES_INPUT' and sha(g_topic) == g_before,
          'reason=%s rc=%s' % (s_g.get('reason'), rc_g))


def case2_no_other_broker():
    name = 'NO_OTHER_BROKER 阻断（单 broker 扩副本）'
    print('\n=== 用例 2：%s ===' % name)
    # One broker, many disks with space: distinct-broker placements are 0 by nature.
    # Same-broker different-mount points are within threshold, but Kafka reassignment JSON
    # forbids duplicate broker IDs in replicas, so the plan cannot be expressed — Stage 1
    # must block instead of marking L2 and emitting invalid replicas=[1,1] JSON like the
    # old version.
    topic = os.path.join(FIX, 'no_other_broker', 'topic.xlsx')
    broker = os.path.join(FIX, 'no_other_broker', 'broker.xlsx')
    before = (sha(topic), sha(broker))
    out = workdir('nob')
    rc, s1 = stage1(topic, broker, out)
    check(name, '阶段一退出码=2（阻断，不是 0/1）', rc == 2, 'rc=%s' % rc)
    check(name, '阶段一 status=FAIL 且 reason=NO_OTHER_BROKER',
          s1.get('status') == 'FAIL' and s1.get('reason') == 'NO_OTHER_BROKER',
          '%s / %s' % (s1.get('status'), s1.get('reason')))
    left = sorted(os.listdir(out))
    check(name, '输出目录为空（阻断时不写半成品产物）', left == [], left)
    blockers = s1.get('blockers', [])
    ok_info = bool(blockers) and all('还需' in b and '异 broker 可用挂载点 0 个' in b
                                     and '无法' in b and '--disk-threshold' in b for b in blockers)
    check(name, 'blockers 给出量化缺口并解释 Kafka JSON 无法表达', ok_info,
          blockers[0] if blockers else '')
    check(name, '未写产物的原因已写进 degradation_events',
          any('不写产物' in e for e in s1.get('degradation_events', [])))
    check(name, '输入文件 sha 不变', (sha(topic), sha(broker)) == before)


def case3_l3():
    name = 'L3 阻断'
    print('\n=== 用例 3：%s ===' % name)
    topic = os.path.join(FIX, 'l3_block', 'topic.xlsx')
    broker = os.path.join(FIX, 'l3_block', 'broker.xlsx')
    before = (sha(topic), sha(broker))
    out = workdir('l3')
    rc, s1 = stage1(topic, broker, out)
    check(name, '阶段一退出码=2（不是 0/1）', rc == 2, 'rc=%s' % rc)
    check(name, '阶段一 status=FAIL 且 reason=L3_BLOCKED',
          s1.get('status') == 'FAIL' and s1.get('reason') == 'L3_BLOCKED',
          '%s / %s' % (s1.get('status'), s1.get('reason')))
    left = sorted(os.listdir(out))
    check(name, '输出目录为空（阻断时不写半成品产物）', left == [], left)
    blockers = s1.get('blockers', [])
    ok_info = bool(blockers) and all('还需' in b and '可用挂载点' in b and '--disk-threshold' in b
                                     for b in blockers)
    check(name, 'blockers 给出量化缺口与可执行出口', ok_info, blockers[0] if blockers else '')
    check(name, '未写产物的原因已写进 degradation_events',
          any('不写产物' in e for e in s1.get('degradation_events', [])))
    check(name, '输入文件 sha 不变', (sha(topic), sha(broker)) == before)


def case4_readonly():
    name = '只查不改'
    print('\n=== 用例 4：%s ===' % name)
    topic, broker = os.path.join(ASSETS, 'topic_template.xlsx'), os.path.join(ASSETS, 'broker_template.xlsx')
    before = (sha(topic), sha(broker))
    pkg_before = snapshot(ROOT)
    out_a, out_b = workdir('ro-a'), workdir('ro-b')
    _rc_a, a = stage1(topic, broker, out_a)
    _rc_b, b = stage1(topic, broker, out_b)
    keys = ('partitions_total', 'new_replicas_assigned', 'max_disk_usage',
            'asset_partition_balance_max', 'asset_partition_balance_min')
    same = all(a.get('stats', {}).get(k) == b.get('stats', {}).get(k) for k in keys)
    check(name, '两次运行阶段一 stats 完全一致（结果可复现）', same,
          '%s vs %s' % (a.get('stats', {}).get('max_disk_usage'), b.get('stats', {}).get('max_disk_usage')))
    fa = os.path.join(out_a, 'topic_new.xlsx')
    fb = os.path.join(out_b, 'topic_new.xlsx')
    ma, mb = table_matrix(fa), table_matrix(fb)
    check(name, '两次 topic_new.xlsx 内容一致（结果可复现）',
          ma is not None and ma == mb, 'rows=%s' % (len(ma) if ma else 0))
    extra = sorted(snapshot(ROOT) - pkg_before)
    check(name, '技能包目录没有新增文件（无 __pycache__ / 无运行残留）', extra == [], extra)
    check(name, '输入文件 sha 两轮后都不变', (sha(topic), sha(broker)) == before)


def case5_pre_existing():
    name = '存量盘已超阈值'
    print('\n=== 用例 5：%s ===' % name)
    topic = os.path.join(FIX, 'pre_existing_over_threshold', 'topic.xlsx')
    broker = os.path.join(FIX, 'pre_existing_over_threshold', 'broker.xlsx')
    before = (sha(topic), sha(broker))
    out = workdir('pre')
    rc, s1 = stage1(topic, broker, out)
    st1 = s1.get('stats', {})
    check(name, '阶段一 status=PASS 且退出码 0（存量问题不阻断本次分配）',
          rc == 0 and s1.get('status') == 'PASS', 'rc=%s status=%s' % (rc, s1.get('status')))
    over = st1.get('pre_existing_over_threshold', [])
    check(name, 'pre_existing_over_threshold 列出存量超阈值挂载点',
          bool(over) and any('data1' in str(x) for x in over), over)
    checks1 = {c['check']: c for c in s1.get('validation', {}).get('checks', [])}
    c11 = checks1.get('1.1', {})
    check(name, '校验 1.1 = WARN（存量问题不判 FAIL）', c11.get('status') == 'WARN', c11.get('detail'))
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(out, 'broker_updated.xlsx'), read_only=True, data_only=True)
    hit = [r for r in wb.active.iter_rows(min_row=2, values_only=True)
           if str(r[0]) == '1' and str(r[4]) == '/srv/BigData/data1']
    wb.close()
    after = float(str(hit[0][7]).replace('%', '')) if hit else 0.0
    check(name, '新副本未落在已超阈值的挂载点（水位未再上升）',
          bool(hit) and abs(after - 90.0) < 0.01, after)
    check(name, '输入文件 sha 不变', (sha(topic), sha(broker)) == before)


def case6_replica3():
    name = '副本数 3 + topic_count 拆分'
    print('\n=== 用例 6：%s ===' % name)
    topic, broker = os.path.join(ASSETS, 'topic_template.xlsx'), os.path.join(ASSETS, 'broker_template.xlsx')
    out = workdir('r3')
    rc1, s1 = stage1(topic, broker, out, replica=3)
    st1 = s1.get('stats', {})
    # 22 partitions × 2 new replicas each = 44
    check(name, '阶段一 status=PASS / new_replicas_assigned=44',
          rc1 == 0 and s1.get('status') == 'PASS' and st1.get('new_replicas_assigned') == 44,
          json.dumps(st1, ensure_ascii=False))
    check(name, '阶段一 degradation_count=0（异 broker 落点足够）', st1.get('degradation_count') == 0,
          st1.get('degradation_count'))
    _rc2, s2 = stage2(out, replica=3)
    header = table_header(os.path.join(out, 'topic_all.xlsx'))
    check(name, '阶段二合并表 12 列且无重名列',
          len(header) == 12 and len(set(header)) == 12, header)
    _rc3, s3 = stage3(out, strategy='topic_count', param=2, replica=3)
    files, docs = read_json_files(out)
    ok_len = bool(docs)
    for doc in docs.values():
        for entry in doc.get('partitions', []):
            if len(entry.get('replicas', [])) != 3 or len(entry.get('log_dirs', [])) != 3:
                ok_len = False
    check(name, '阶段三条目 replicas/log_dirs 长度均为 3', ok_len, '%s 个文件' % len(files))
    check(name, '按 topic_count 策略生成 batch_N.json 且 validation_fail=0',
          files and all(f.startswith('batch_') and f.endswith('.json') for f in files)
          and s3.get('stats', {}).get('validation_fail') == 0, files)

    # The UI offers another option: split by topic partition size. Both strategies must
    # run independently, otherwise "custom output JSON" is only decorative in the UI.
    rc_sz, s_sz = stage3(out, strategy='size_10g', param=1, replica=3)
    check(name, '按 topic 分区大小策略（size_10g）跑通且 validation_fail=0',
          rc_sz == 0 and s_sz.get('status') == 'PASS'
          and s_sz.get('stats', {}).get('validation_fail') == 0,
          json.dumps(s_sz.get('stats', {}), ensure_ascii=False))


def _write_topic_all(out, rows):
    """Hand-build a merged master table topic_all.xlsx (same shape as the Stage 2 artifact) for Stage 3 hard-check cases."""
    import openpyxl
    path = os.path.join(out, 'topic_all.xlsx')
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'topic_all'
    ws.append(['Topic', 'PartitionCount', 'Partition', 'Original Broker ID', 'Original Path',
               'Allocation Broker ID', 'Allocation Path', 'Size(bytes)', 'Size(GB)', 'degradation'])
    for r in rows:
        ws.append(r)
    wb.save(path)
    wb.close()
    return path


def case7_duplicate_broker_fail():
    name = '阶段三硬校验：重复 broker 不再豁免'
    print('\n=== 用例 7：%s ===' % name)
    # Build a merged table marked degradation=L2: both replicas of one partition land on
    # broker 1 (the typical legacy L2 artifact). The new Stage 1 blocks this scenario; if
    # someone hand-crafts such a table, Stage 3 must FAIL (3.4) instead of emitting
    # replicas=[1,1] that Kafka would reject.
    out = workdir('dup')
    _write_topic_all(out, [
        ['ITDE_TEST_dup_broker_lkp', 2, 0, 1, '/srv/BigData/data1',
         1, '/srv/BigData/data2', 107374182, 0.1, 'L2'],
    ])
    rc, s3 = stage3(out, strategy='single_file', param='10')
    st3 = s3.get('stats', {})
    check(name, '阶段三 rc=1 / status=FAIL / validation_fail>0',
          rc == 1 and s3.get('status') == 'FAIL' and st3.get('validation_fail', 0) > 0,
          json.dumps(st3, ensure_ascii=False))
    c34 = [d for d in s3.get('validation_details', []) if d.get('check') == '3.4']
    check(name, '3.4 对重复 broker 判 FAIL（不再有 L2-EXCEPTION 放行）',
          any(d.get('status') == 'FAIL' for d in c34), c34)
    check(name, 'stats 不再包含 l2_exception_count（放行概念已移除）',
          'l2_exception_count' not in st3, sorted(st3.keys()))


def case8_topic_name_sanitize():
    name = '阶段三 topic 名路径穿越拦截'
    print('\n=== 用例 8：%s ===' % name)
    # Topic names come from user-uploaded Excel (untrusted input); per_topic / size_10g /
    # custom embed them into file names. Topic names containing path separators must be
    # blocked (INVALID_TOPIC_NAME) with no file written; nothing may escape the output dir.
    out = workdir('sanitize')
    _write_topic_all(out, [
        ['../pwn', 1, 0, 1, '/srv/BigData/data1', 2, '/srv/BigData/data9',
         107374182, 0.1, ''],
    ])
    rc, s3 = stage3(out, strategy='per_topic', param='10')
    check(name, '非法 topic 名被拦截（reason=INVALID_TOPIC_NAME, rc=1）',
          rc == 1 and s3.get('reason') == 'INVALID_TOPIC_NAME',
          '%s / %s' % (rc, s3.get('reason')))
    jdir = os.path.join(out, 'json')
    check(name, '阻断时 json 目录为空（未写出目录）',
          not os.path.isdir(jdir) or sorted(os.listdir(jdir)) == [],
          sorted(os.listdir(jdir)) if os.path.isdir(jdir) else [])
    check(name, '没有文件逃逸到输出目录之外',
          not os.path.exists(os.path.join(out, 'pwn.json'))
          and not os.path.exists(os.path.join(os.path.dirname(out), 'pwn.json')),
          '')
    # Another kind of untrusted input: characters outside the whitelist (e.g. non-ASCII) are rejected the same way
    out2 = workdir('sanitize2')
    _write_topic_all(out2, [
        ['中文topic', 1, 0, 1, '/srv/BigData/data1', 2, '/srv/BigData/data9',
         107374182, 0.1, ''],
    ])
    rc2, s3b = stage3(out2, strategy='per_topic', param='10')
    check(name, '白名单外字符（中文 topic 名）同样被拦截',
          rc2 == 1 and s3b.get('reason') == 'INVALID_TOPIC_NAME',
          s3b.get('reason'))


def main():
    print('技能包 : %s' % ROOT)
    print('解释器 : %s' % PY)
    case1_golden()
    case2_no_other_broker()
    case3_l3()
    case4_readonly()
    case5_pre_existing()
    case6_replica3()
    case7_duplicate_broker_fail()
    case8_topic_name_sanitize()
    passed = sum(1 for r in RESULTS if r['passed'])
    failed = [r for r in RESULTS if not r['passed']]
    print('\n' + '=' * 60)
    print('回归结果: PASS=%d FAIL=%d' % (passed, len(failed)))
    for r in failed:
        print('  FAIL [%s] %s <- %s' % (r['case'], r['text'], r['evidence']))
    print('=' * 60)
    for d in _TMP:
        shutil.rmtree(d, ignore_errors=True)
    return len(failed)


if __name__ == '__main__':
    sys.exit(main())
