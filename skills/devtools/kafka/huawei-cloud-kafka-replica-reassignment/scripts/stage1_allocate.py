# AI-generated
#!/usr/bin/env python3
# AI-generated
# AI-generated
"""
Stage 1: Allocate new replicas for Kafka partitions.

Inputs: topic.xlsx (current partitions), broker.xlsx (disk usage)
Artifacts: topic.xlsx (original replica table), topic_new.xlsx (new replicas),
           broker_updated.xlsx (refreshed disk table)

Placement & blocking (see references/stage1-allocate.md):
    L0  replicas land on different brokers, projected usage within threshold (normal path)
    NO_OTHER_BROKER  only same-broker different-mount points are within threshold: Kafka
                     reassignment JSON forbids duplicate broker IDs in one partition's
                     replicas, so "same broker, different mount point" cannot be expressed
                     -- block: no artifacts, status=FAIL, exit code 2
    L3  no mount point can host the replica within threshold -- block: no artifacts,
        status=FAIL, exit code 2

Exit codes: 0 = PASS / 1 = stage validation has FAIL / 2 = blocked (NO_OTHER_BROKER
no distinct broker / L3_BLOCKED insufficient capacity, add capacity or brokers first).
stdout carries only the stage-result JSON; the human-readable four-line stage report goes
to stderr so it never pollutes the JSON.
"""

import argparse
import json
import os
import sys

from collections import defaultdict
from copy import deepcopy

import openpyxl

# The skill package may be read-only after install, and no compile cache belongs in it:
# disable bytecode writing so execution never creates __pycache__ under scripts/.
sys.dont_write_bytecode = True

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import tabio   # noqa: E402  same-dir module; must import after sys.path fix

EXIT_PASS = 0
EXIT_VALIDATION_FAIL = 1
EXIT_L3_BLOCKED = 2   # shared by all blocking reasons: L3_BLOCKED (capacity) and NO_OTHER_BROKER (no distinct broker)

STAGE_NAME = '阶段一'


def read_topic_xlsx(path):
    """Read topic.xlsx -> list of dicts with keys:
    Topic, PartitionCount, Partition, Broker ID, Path, Size(bytes), Size(GB)
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    headers = [str(h).strip() if h else '' for h in rows[0]]
    # Normalize header names (remove extra spaces)
    header_map = {}
    for i, h in enumerate(headers):
        hl = h.lower().replace(' ', '').replace('_', '')
        if 'topic' in hl and 'count' not in hl and 'partition' not in hl:
            header_map['Topic'] = i
        elif 'partitioncount' in hl:
            header_map['PartitionCount'] = i
        elif hl == 'partition':
            header_map['Partition'] = i
        elif 'brokerid' in hl or (hl == 'brokerid'):
            header_map['Broker ID'] = i
        elif hl == 'path':
            header_map['Path'] = i
        elif 'size(bytes)' in hl or 'sizebytes' in hl:
            header_map['Size(bytes)'] = i
        elif 'size(gb)' in hl or 'sizegb' in hl:
            header_map['Size(GB)'] = i

    records = []
    for row in rows[1:]:
        if row is None or all(c is None for c in row):
            continue
        rec = {}
        for key, idx in header_map.items():
            val = row[idx] if idx < len(row) else None
            if key in ('PartitionCount', 'Partition', 'Broker ID', 'Size(bytes)'):
                try:
                    val = int(val) if val is not None else 0
                except (ValueError, TypeError):
                    val = 0
            elif key == 'Size(GB)':
                try:
                    val = float(val) if val is not None else 0.0
                except (ValueError, TypeError):
                    val = 0.0
            rec[key] = val
        if rec.get('Topic'):
            records.append(rec)
    return records


def read_broker_xlsx(path):
    """Read broker.xlsx -> list of dicts; keys follow the broker-table column names
    (Broker ID, Host, Disk summary, device, mount point, used GB, total GB, usage %).
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    headers = [str(h).strip() if h else '' for h in rows[0]]
    header_map = {}
    for i, h in enumerate(headers):
        hl = h.lower().replace(' ', '').replace('_', '').replace('（', '(').replace('）', ')')
        if 'brokerid' in hl:
            header_map['Broker ID'] = i
        elif hl == 'host':
            header_map['Host'] = i
        elif 'disk' in hl and 'used' in hl:
            header_map['Disk (Used | Total)'] = i
        elif hl == '分区' or hl == 'mount' or 'device' in hl:
            header_map['分区'] = i
        elif '挂载点' in hl or hl == 'path' or 'mountpoint' in hl:
            header_map['挂载点'] = i
        elif '已使用' in hl or 'used(gb)' in hl:
            header_map['已使用（GB）'] = i
        elif '磁盘大小' in hl or 'total(gb)' in hl or 'disksize' in hl:
            header_map['磁盘大小（GB）'] = i
        elif '使用率' in hl or 'usage' in hl:
            header_map['磁盘使用率'] = i

    records = []
    for row in rows[1:]:
        if row is None or all(c is None for c in row):
            continue
        rec = {}
        for key, idx in header_map.items():
            val = row[idx] if idx < len(row) else None
            if key == 'Broker ID':
                try:
                    val = int(val) if val is not None else 0
                except (ValueError, TypeError):
                    val = 0
            elif key in ('已使用（GB）', '磁盘大小（GB）'):
                try:
                    val = float(val) if val is not None else 0.0
                except (ValueError, TypeError):
                    val = 0.0
            elif key == '磁盘使用率':
                if val is None:
                    val = 0.0
                else:
                    s = str(val).replace('%', '').strip()
                    try:
                        val = float(s)
                    except (ValueError, TypeError):
                        val = 0.0
            rec[key] = val
        if rec.get('Broker ID'):
            records.append(rec)
    return records


def group_topics(topic_records):
    """Group topic records by topic name, preserving partition order."""
    topics = defaultdict(list)
    for rec in topic_records:
        topics[rec['Topic']].append(rec)
    for t in topics:
        topics[t].sort(key=lambda r: r['Partition'])
    return topics


def find_duplicate_partitions(topic_records):
    """Find records where the same (Topic, Partition) appears more than once.

    This tool handles "single-replica current state -> N replicas" expansion; if the input
    table already has multiple rows for one partition, the current state is not
    single-replica and continuing would expand twice, so we must stop and ask.
    The comparison key normalizes leading/trailing whitespace of the Topic so duplicates
    like "abc" and "abc " (invisible to the eye) cannot slip past the gate and carry a
    545-row corrupted table into later stages.
    """
    seen = defaultdict(int)
    for rec in topic_records:
        seen[(str(rec['Topic']).strip(), rec['Partition'])] += 1
    return [k for k, v in seen.items() if v > 1]


def survey_pre_existing_over_threshold(broker_records, disk_threshold):
    """List input broker mount points that are already over threshold.

    The threshold constrains the projected usage of NEW replicas; it cannot roll back
    existing watermarks. Listing pre-existing over-threshold mount points explicitly
    avoids a silent relaxation like "status=PASS but max_disk_usage above threshold".
    """
    over = []
    for rec in broker_records:
        usage = rec.get('磁盘使用率', 0.0)
        if usage > disk_threshold:
            over.append('broker %s %s: %.2f%%' % (rec['Broker ID'], rec.get('挂载点', ''), usage))
    return over


def get_broker_mountpoints(broker_records):
    """Build a lookup: (broker_id, mount_path) -> broker record dict."""
    lookup = {}
    for rec in broker_records:
        lookup[(rec['Broker ID'], rec.get('挂载点', ''))] = rec
    return lookup


def get_all_broker_ids(broker_records):
    """Get sorted unique broker IDs."""
    return sorted(set(r['Broker ID'] for r in broker_records))


def get_mountpoints_for_broker(broker_records, broker_id):
    """Get all mountpoint records for a given broker, sorted by usage."""
    mounts = [r for r in broker_records if r['Broker ID'] == broker_id]
    mounts.sort(key=lambda r: r.get('磁盘使用率', 0.0))
    return mounts


FAIRNESS_BAND = 1.0   # within 1 pt of usage the disks are "comparable"; judge load fairness


def _build_candidates(broker_records, used_pairs, size_gb, same_broker_ids,
                      alloc_counts=None, band=FAIRNESS_BAND):
    """List usable mount-point candidates, ascending by "disk tier -> replicas already
    accepted -> usage" (stable sort, reproducible).

    (broker, mount point) pairs already used by another replica of this partition are
    excluded outright: placing two replicas on the same mount point would break the
    "different mount point" premise and double-count capacity on one disk.

    Why not sort purely by usage: partition sizes (0.1 GB scale) barely change usage
    relative to disks (900 GB scale), so a strict usage sort would keep hitting the same
    lowest-watermark disk — a real test placed 15 of 22 new replicas onto one broker and
    left the other 14 brokers untouched. Kafka's replica distribution directly decides
    failure blast radius and read/write balance, so mount points within 1 percentage
    point of the lowest watermark are treated as the same tier (comparable disks); within
    a tier, prefer the broker that accepted the fewest replicas in this run; within a
    broker, sort by usage. Picking the lightest-loaded mount inside the tier keeps the
    "lowest usage first" intent without piling replicas onto a few brokers.
    """
    candidates = []
    for rec in broker_records:
        bid = rec['Broker ID']
        mnt = rec.get('挂载点', '')
        if (bid, mnt) in used_pairs:
            continue
        used_gb = rec.get('已使用（GB）', 0.0)
        total_gb = rec.get('磁盘大小（GB）', 0.0)
        proj_used = used_gb + size_gb
        proj_usage = (proj_used / total_gb * 100) if total_gb > 0 else 100.0
        candidates.append({
            'bid': bid,
            'mnt': mnt,
            'usage': rec.get('磁盘使用率', 0.0),
            'proj_usage': proj_usage,
            'total': total_gb,
            'free': max(total_gb - used_gb, 0.0),
            'same_broker': bid in same_broker_ids,
            'rec': rec,
        })
    if not candidates:
        return candidates
    counts = alloc_counts or {}
    lowest = min(c['usage'] for c in candidates)
    for c in candidates:
        c['disk_tier'] = 0 if c['usage'] <= lowest + band else 1
        c['load'] = counts.get(c['bid'], 0)
    candidates.sort(key=lambda c: (c['disk_tier'], c['load'], c['usage']))
    return candidates


def describe_l3_block(topic_name, partition_id, size_gb, slots_left, candidates, disk_threshold):
    """Give an actionable, quantified gap for a blocked replica (mount points short / GB short).

    L3 means "cannot be done within the threshold", not "looks ugly". Reporting only a
    failure gives nothing to act on, so the gap is computed here: the user can add mount
    points or relax the threshold, then re-run.
    """
    usable = [c for c in candidates if c['proj_usage'] <= disk_threshold]
    distinct_usable = [c for c in usable if not c['same_broker']]
    same_broker_usable = [c for c in usable if c['same_broker']]
    mounts_short = max(0, slots_left - len(distinct_usable))
    need_gb = size_gb * slots_left
    have_gb = sum(sorted((c['free'] for c in usable), reverse=True)[:slots_left])
    space_short = max(0.0, need_gb - have_gb)
    return (
        'Topic=%s, Partition=%s: 还需 %d 个副本（每个 %.2fGB）；阈值 %s%% 内可用挂载点 %d 个'
        '（异 broker %d 个，同 broker %d 个）。'
        '只增加同 broker 落点对该副本无济于事（会重复 broker ID，无法表达为 Kafka reassignment JSON）；'
        '需让至少 %d 个异 broker 挂载点进入阈值：新增异 broker 挂载点，或释放其他 broker 容量共 %.2fGB，'
        '或提高 --disk-threshold。'
        % (topic_name, partition_id, slots_left, size_gb, disk_threshold,
           len(usable), len(distinct_usable), len(same_broker_usable), mounts_short, space_short)
    )


def describe_no_other_broker_block(topic_name, partition_id, size_gb, slots_left,
                                   usable, distinct_usable, disk_threshold):
    """Give an actionable, quantified gap for a replica with no distinct broker to land on.

    Same-broker different-mount points may be within threshold, but standard Kafka
    reassignment JSON forbids duplicate broker IDs in one partition's replicas (log_dirs
    can only name directories for existing replicas), so this is not "a degradation we can
    do" but "inexpressible, must block". Quantify: how many distinct-broker placements are
    missing, how many same-broker mount points are wasted, how to fix it.
    """
    same = len(usable) - len(distinct_usable)
    return (
        'Topic=%s, Partition=%s: 还需 %d 个副本（每个 %.2fGB）；阈值 %s%% 内可用挂载点 %d 个，'
        '但异 broker 可用挂载点 0 个（同 broker 不同挂载点 %d 个——同一 partition 的 replicas '
        '不允许重复 broker ID，Kafka reassignment JSON 无法表达这种方案）。'
        '需新增/扩容一个异 broker 挂载点，或释放其他 broker 的容量，或提高 --disk-threshold。'
        % (topic_name, partition_id, slots_left, size_gb, disk_threshold,
           len(usable), same)
    )


def allocate_replicas(topic_records, broker_records, replica_count, disk_threshold):
    """Allocate new replicas for every partition.

    Returns a dict (keys differ from the old version; callers must unpack via return keys):
        new_records            topic_new.xlsx records (degradation always empty: only L0 admitted)
        updated_brokers        broker table with refreshed disk usage
        degradation_log        blocking events (human-readable)
        no_other_broker_blockers  no-distinct-broker blocks; non-empty means stop, no artifacts
        l3_blockers            L3 capacity blocks; non-empty means stop, no artifacts
        broker_partition_count partitions hosted per broker (for the 1.6 balance check)
        new_replica_per_broker new replicas accepted per broker in this run

    Order: L0 (distinct broker) -> block when no distinct broker is available. Two kinds:
        NO_OTHER_BROKER: only same-broker different-mount points are within threshold. The
            old version marked L2 and continued to emit JSON like replicas=[1,1], but Kafka
            reassignment JSON forbids duplicate broker IDs (log_dirs can only name
            directories for existing replicas), so such JSON fails on submission — an
            undeliverable plan; block instead of degrading.
        L3_BLOCKED: no mount point at all is within threshold (insufficient capacity).
    There is deliberately no "force it on anyway" fallback: silently relaxing the disk
    threshold just pushes capacity risk into production, while "no artifacts + exit code 2"
    makes the pipeline and the operator stop and fix it immediately.
    """
    topics = group_topics(topic_records)
    updated_brokers = deepcopy(broker_records)

    broker_partition_count = defaultdict(int)
    for rec in topic_records:
        broker_partition_count[rec['Broker ID']] += 1

    new_records = []
    degradation_log = []
    no_other_broker_blockers = []
    l3_blockers = []
    alloc_counts = defaultdict(int)   # new replicas accepted per broker this run, fed to candidate sort for fairness

    for topic_name in sorted(topics.keys()):
        for part_rec in topics[topic_name]:
            partition_id = part_rec['Partition']
            size_gb = part_rec.get('Size(GB)', 0.0) or 0.0
            size_bytes = part_rec.get('Size(bytes)', 0)
            origin_broker = part_rec['Broker ID']
            used_brokers = {origin_broker}
            used_pairs = {(origin_broker, part_rec.get('Path', ''))}
            slots = replica_count - 1

            for slot in range(slots):
                candidates = _build_candidates(updated_brokers, used_pairs, size_gb, used_brokers,
                                                alloc_counts=alloc_counts)

                selected = None
                for c in candidates:
                    if not c['same_broker'] and c['proj_usage'] <= disk_threshold:
                        selected = (c, 'L0')
                        break

                if selected is None:
                    # No compliant distinct-broker placement -> block, split by cause:
                    # - no usable mount point at all within threshold -> L3 (capacity)
                    # - only same-broker different-mount points within threshold -> NO_OTHER_BROKER
                    #   Same-broker different-mount cannot be expressed in standard Kafka
                    #   reassignment JSON (duplicate broker IDs are forbidden), and the old L2
                    #   output replicas=[1,1] would be rejected by Kafka; must block.
                    usable = [c for c in candidates if c['proj_usage'] <= disk_threshold]
                    distinct_usable = [c for c in usable if not c['same_broker']]
                    if not usable:
                        l3_blockers.append(describe_l3_block(
                            topic_name, partition_id, size_gb, slots - slot,
                            candidates, disk_threshold))
                        degradation_log.append(
                            'Topic=%s, Partition=%s: L3 阻断 - 阈值 %s%% 内无任何可用挂载点，'
                            '已停止本阶段且不写产物' % (topic_name, partition_id, disk_threshold))
                    else:
                        no_other_broker_blockers.append(describe_no_other_broker_block(
                            topic_name, partition_id, size_gb, slots - slot,
                            usable, distinct_usable, disk_threshold))
                        degradation_log.append(
                            'Topic=%s, Partition=%s: NO_OTHER_BROKER 阻断 - 异 broker 可用挂载点 0 个'
                            '（同 broker 不同挂载点 %d 个，无法在 Kafka reassignment JSON 中表达），'
                            '已停止本阶段且不写产物'
                            % (topic_name, partition_id, len(usable) - len(distinct_usable)))
                    break

                c, deg = selected
                bid, mnt, rec = c['bid'], c['mnt'], c['rec']
                used_brokers.add(bid)
                used_pairs.add((bid, mnt))
                broker_partition_count[bid] += 1
                alloc_counts[bid] += 1

                rec['已使用（GB）'] = rec.get('已使用（GB）', 0.0) + size_gb
                total_gb = rec.get('磁盘大小（GB）', 0.0)
                if total_gb > 0:
                    rec['磁盘使用率'] = round(rec['已使用（GB）'] / total_gb * 100, 2)

                new_records.append({
                    'Topic': topic_name,
                    'PartitionCount': part_rec['PartitionCount'],
                    'Partition': partition_id,
                    'Broker ID': bid,
                    'Path': mnt,
                    'Size(bytes)': size_bytes,
                    'Size(GB)': size_gb,
                    'degradation': '' if deg == 'L0' else deg,
                })

    return {
        'new_records': new_records,
        'updated_brokers': updated_brokers,
        'degradation_log': degradation_log,
        'no_other_broker_blockers': no_other_broker_blockers,
        'l3_blockers': l3_blockers,
        'broker_partition_count': dict(broker_partition_count),
        'new_replica_per_broker': dict(alloc_counts),
    }


NEW_HEADERS = ['Topic', 'PartitionCount', 'Partition', 'Broker ID', 'Path',
               'Size(bytes)', 'Size(GB)', 'degradation']
BROKER_HEADERS = ['Broker ID', 'Host', 'Disk (Used | Total)', '分区', '挂载点',
                  '已使用（GB）', '磁盘大小（GB）', '磁盘使用率']


MAIN_HEADERS = ['Topic', 'PartitionCount', 'Partition', 'Broker ID', 'Path',
                'Size(bytes)', 'Size(GB)']


def write_new_table(new_records, output_path):
    """Write topic_new.xlsx (new replica table; trailing degradation column carries markers).

    Same xlsx format as the inputs: this table is usually checked against Excel and handed
    to others; TSV would need a manual conversion and Chinese headers often garble in Excel.
    """
    rows = [[rec.get(h, '') for h in NEW_HEADERS] for rec in new_records]
    tabio.write_table(output_path, NEW_HEADERS, rows)


def write_main_table(topic_records, output_path):
    """Write topic.xlsx (original replica table, consumed by Stage 2)."""
    rows = [[rec.get(h, '') for h in MAIN_HEADERS] for rec in topic_records]
    tabio.write_table(output_path, MAIN_HEADERS, rows)


def write_broker_xlsx(broker_records, output_path):
    """Write the refreshed broker table; usage is written as an xx.xx% string to match the
    human-readable table."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Broker'
    ws.append(BROKER_HEADERS)
    for rec in broker_records:
        ws.append([
            rec.get('Broker ID', ''),
            rec.get('Host', ''),
            rec.get('Disk (Used | Total)', ''),
            rec.get('分区', ''),
            rec.get('挂载点', ''),
            rec.get('已使用（GB）', 0.0),
            rec.get('磁盘大小（GB）', 0.0),
            '%.2f%%' % rec.get('磁盘使用率', 0.0),
        ])
    wb.save(output_path)


def run(topic_path, broker_path, output_dir, replica_count=2, disk_threshold=70):
    """Stage 1 main flow: read -> allocate -> write artifacts.

    On an L3/NO_OTHER_BROKER block **no artifact is written** (a half-finished table must
    not be mistaken by later stages for a usable plan); only a FAIL result is returned.
    """
    os.makedirs(output_dir, exist_ok=True)

    topic_records = read_topic_xlsx(topic_path)
    broker_records = read_broker_xlsx(broker_path)
    params = {'replica_count': replica_count, 'disk_threshold': disk_threshold}

    if not topic_records:
        return {'status': 'FAIL', 'stage': 1, 'reason': 'BAD_INPUT',
                'message': 'topic 表无有效记录：%s' % topic_path,
                'inputs': [topic_path, broker_path], 'outputs': [], 'params': params,
                'degradation_events': [], 'degradation_count': 0, 'stats': {}}
    if not broker_records:
        return {'status': 'FAIL', 'stage': 1, 'reason': 'BAD_INPUT',
                'message': 'broker 表无有效记录：%s' % broker_path,
                'inputs': [topic_path, broker_path], 'outputs': [], 'params': params,
                'degradation_events': [], 'degradation_count': 0, 'stats': {}}
    if replica_count < 2:
        return {'status': 'FAIL', 'stage': 1, 'reason': 'BAD_PARAM',
                'message': 'replica_count 必须 >= 2（当前 %s）' % replica_count,
                'inputs': [topic_path, broker_path], 'outputs': [], 'params': params,
                'degradation_events': [], 'degradation_count': 0, 'stats': {}}

    dups = find_duplicate_partitions(topic_records)
    if dups:
        return {'status': 'FAIL', 'stage': 1, 'reason': 'BAD_INPUT',
                'message': '输入 topic 表同一分区出现多行（现状非单副本）：%s；'
                           '本工具只处理「单副本现状 → N 副本」扩容，请先确认现状。' % dups[:5],
                'inputs': [topic_path, broker_path], 'outputs': [], 'params': params,
                'degradation_events': [], 'degradation_count': 0, 'stats': {}}

    # The artifact and the input table share the name topic.xlsx: if --output-dir points at
    # the input directory, writing would overwrite the user's current-state data, irreversibly.
    # Failing here is better than overwriting the input.
    main_path = os.path.join(output_dir, tabio.MAIN_NAME)
    if tabio.same_file(main_path, topic_path):
        return {'status': 'FAIL', 'stage': 1, 'reason': 'OUTPUT_OVERWRITES_INPUT',
                'message': '输出目录与输入 topic.xlsx 指向同一路径（%s）：产物会覆盖现状数据表，'
                           '请把 --output-dir 指向单独目录后重跑。' % main_path,
                'inputs': [topic_path, broker_path], 'outputs': [], 'params': params,
                'degradation_events': [], 'degradation_count': 0, 'stats': {}}

    alloc = allocate_replicas(topic_records, broker_records, replica_count, disk_threshold)
    new_records = alloc['new_records']
    updated_brokers = alloc['updated_brokers']
    degradation_log = alloc['degradation_log']
    no_other_broker_blockers = alloc['no_other_broker_blockers']
    l3_blockers = alloc['l3_blockers']

    base_stats = {
        'topics_processed': len(set(r['Topic'] for r in topic_records)),
        'partitions_total': len(topic_records),
        'new_replicas_assigned': len(new_records),
        'degradation_count': len(degradation_log),
        # This version no longer produces L2 degradation (same-broker different-mount cannot
        # be expressed as Kafka reassignment JSON); the key is kept for downstream parsing and
        # is always 0.
        'l2_degradation_count': 0,
    }

    if no_other_broker_blockers or l3_blockers:
        reason = 'NO_OTHER_BROKER' if no_other_broker_blockers else 'L3_BLOCKED'
        blocked = no_other_broker_blockers + l3_blockers
        if reason == 'NO_OTHER_BROKER':
            message = ('无可用异 broker：阈值 %s%% 内只剩同 broker 的不同挂载点，'
                       '无法生成 Kafka reassignment JSON（同一 partition 的 replicas 不允许'
                       '重复 broker ID），本阶段未写任何产物。' % disk_threshold)
            next_action = '新增 broker，或在其他 broker 上释放容量/提高 --disk-threshold 后重跑阶段一'
        else:
            message = ('挂载点/容量不足：阈值 %s%% 内无法为 %d 个副本找到落点，'
                       '本阶段未写任何产物（不再静默强行分配）。'
                       % (disk_threshold, len(l3_blockers)))
            next_action = '扩容挂载点/增加 broker，或提高 --disk-threshold 后重跑阶段一'
        return {
            'status': 'FAIL',
            'stage': 1,
            'reason': reason,
            'message': message,
            'inputs': [topic_path, broker_path],
            'outputs': [],
            'params': params,
            'blockers': blocked,
            'degradation_events': degradation_log,
            'degradation_count': len(degradation_log),
            'stats': dict(base_stats,
                          blocked_replicas=len(blocked),
                          no_other_broker_blocked_replicas=len(no_other_broker_blockers),
                          l3_blocked_replicas=len(l3_blockers)),
            'next_action': next_action,
        }

    new_path = os.path.join(output_dir, tabio.NEW_NAME)
    broker_updated_path = os.path.join(output_dir, tabio.BROKER_UPDATED_NAME)
    write_main_table(topic_records, main_path)
    write_new_table(new_records, new_path)
    write_broker_xlsx(updated_brokers, broker_updated_path)

    max_usage = max((r.get('磁盘使用率', 0.0) for r in updated_brokers), default=0.0)
    counts = list(alloc['broker_partition_count'].values())
    over = survey_pre_existing_over_threshold(broker_records, disk_threshold)

    return {
        'status': 'PASS',
        'stage': 1,
        'inputs': [topic_path, broker_path],
        'outputs': [main_path, new_path, broker_updated_path],
        'params': params,
        'degradation_events': degradation_log,
        'degradation_count': len(degradation_log),
        'stats': dict(base_stats,
                      max_disk_usage=round(max_usage, 2),
                      pre_existing_over_threshold=over,
                      asset_partition_balance_max=max(counts) if counts else 0,
                      asset_partition_balance_min=min(counts) if counts else 0),
    }


def _stage_report(stage_no, inputs, outputs, validation):
    """Four-line stage report (human-readable): stage / inputs / outputs / validation. The Web
    UI and the CLI share the same wording."""
    inputs = inputs or []
    outputs = outputs or []
    if validation:
        verdict = 'PASS' if validation.get('pass') else 'FAIL'
        summary = validation.get('summary', '') or ''
        conclusions = [l.strip() for l in summary.splitlines() if l.strip().startswith('结论')]
        detail = conclusions[-1] if conclusions else (summary.splitlines()[-1].strip() if summary else '')
        vtext = '%s  %s' % (verdict, detail)
    else:
        vtext = '未执行（本阶段未进入校验，原因见 status/reason）'
    return '阶段：%s\n输入：%s\n产物：%s\n校验：%s' % (
        stage_no,
        ' | '.join(str(x) for x in inputs) or '无',
        ' | '.join(str(x) for x in outputs) or '无',
        vtext,
    )


def main(argv=None):
    parser = argparse.ArgumentParser(description='Stage 1: 为 Kafka 分区分配新副本')
    parser.add_argument('--topic', required=True, help='topic.xlsx 路径')
    parser.add_argument('--broker', required=True, help='broker.xlsx 路径')
    parser.add_argument('--output-dir', required=True, help='产物输出目录')
    parser.add_argument('--replica-count', type=int, default=2, help='目标副本数（默认 2）')
    parser.add_argument('--disk-threshold', type=float, default=70,
                        help='磁盘使用率阈值 %%（默认 70）')
    args = parser.parse_args(argv)

    result = run(args.topic, args.broker, args.output_dir,
                 args.replica_count, args.disk_threshold)

    # The stage gate must live on both the Web UI and the command line: validating only in
    # the UI means headless execution has no gate at all.
    if result.get('status') == 'PASS':
        try:
            from validate import validate_stage1, format_checks

            checks = validate_stage1(result, args.output_dir, args.disk_threshold)
            text, passed = format_checks(STAGE_NAME, checks)
            result['validation'] = {
                'checks': [{'check': c[0], 'status': c[1], 'detail': c[2]} for c in checks],
                'summary': text,
                'pass': passed,
            }
            if not passed:
                result['status'] = 'FAIL'
                result['reason'] = 'VALIDATION_FAIL'
        except Exception as exc:  # an unavailable validator must also block, never pass silently
            result['status'] = 'FAIL'
            result['reason'] = 'VALIDATOR_UNAVAILABLE'
            result['validation'] = {'checks': [], 'summary': '校验模块不可用: %s' % exc,
                                    'pass': False}

    exit_code = EXIT_PASS
    if result.get('status') != 'PASS':
        exit_code = EXIT_L3_BLOCKED if result.get('reason') in ('L3_BLOCKED', 'NO_OTHER_BROKER') else EXIT_VALIDATION_FAIL

    # stdout carries only JSON (upstream parses the stage result); the human-readable
    # four-line report goes to stderr
    print(json.dumps(result, ensure_ascii=False, indent=2))
    print(_stage_report('阶段一', result.get('inputs'), result.get('outputs'),
                        result.get('validation')), file=sys.stderr)
    return exit_code


if __name__ == '__main__':
    sys.exit(main())